"""
このテストが守るもの: ラベル比較Excelの行が、Status列（A Only／B Only／Different／Same）
の値ごとに全区分淡色の配色セット（ピンク／青／黄／緑）で色分けされること。

受入条件（2026-10-14、ユーザー指定。HostPL-comparer・DXF-label-comparison と同じ配色）:
    A Only=薄いピンク／B Only=薄い青／Different=薄い黄／Same=薄い緑。
    従来はSame行に書式が無く（無色）、A Onlyが青・B Onlyが緑だった。
"""
import io
from types import SimpleNamespace

import ezdxf
import pytest
from openpyxl import load_workbook

from utils.compare_labels import STATUS_COLORS, compare_labels_multi


def _write_dxf(path, labels):
    doc = ezdxf.new()
    msp = doc.modelspace()
    for i, label in enumerate(labels):
        msp.add_text(label, dxfattribs={'insert': (10 * i, 0)})
    doc.saveas(path)


@pytest.fixture
def excel_bytes(tmp_path):
    a, b = tmp_path / 'a.dxf', tmp_path / 'b.dxf'
    _write_dxf(a, ['R1', 'R1', 'C1', 'U1'])   # R1:2 / C1:1 / U1:1
    _write_dxf(b, ['R1', 'R1', 'C1', 'C1', 'L1'])  # R1:2 / C1:2 / L1:1
    pair = (SimpleNamespace(name='a.dxf'), SimpleNamespace(name='b.dxf'), str(a), str(b), 'P')
    return compare_labels_multi([pair])


def test_status_rows_use_light_color_set(excel_bytes):
    ws = load_workbook(io.BytesIO(excel_bytes))['P']
    statuses = {ws.cell(r, 4).value for r in range(2, ws.max_row + 1)}
    assert statuses == {'A Only', 'B Only', 'Different', 'Same'}

    rules = {}
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            status = rule.formula[0].split('"')[1]
            rules[status] = rule.dxf.fill.bgColor.rgb
    assert set(rules) == set(STATUS_COLORS)
    for status, color in STATUS_COLORS.items():
        assert rules[status].endswith(color['bg_color'][1:])
    assert STATUS_COLORS['A Only']['bg_color'] == '#FEF1F5'
    assert STATUS_COLORS['B Only']['bg_color'] == '#EAF1FB'
    assert STATUS_COLORS['Different']['bg_color'] == '#FFF6D5'
    assert STATUS_COLORS['Same']['bg_color'] == '#EAF7EC'
